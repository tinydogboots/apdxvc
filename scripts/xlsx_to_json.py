"""
Convert audit XLSX files → data/kumu_blueprint.json

Workbook layout (one sheet "Task Database", 13 columns)
────────────────────────────────────────────────────────
  Task ID, Tier, Category, Sub-heading, Source Sentence, Task,
  Lead, Track, Recommendation Only,
  Connection Type, Connected Task IDs, Connection Pending Confirmation, Notes

Connections are inline: each row's "Connection Type" + "Connected Task IDs"
defines outbound edges.  Direction follows project flow — when row X has
"Dependency" → "Y", we emit an edge Y → X (Y enables / is a prerequisite of X).
Synergy edges are de-duped per unordered pair.

Lead normalisation
──────────────────
  Compound values (e.g. "Client / IT", "Facilities / Third-party") are reduced
  to one of {consultant, client, third_party, tbd} for filtering.  The
  original raw string is preserved in `attributes.lead_label`.

Usage
─────
    python scripts/xlsx_to_json.py                     # reads data/*.xlsx
    python scripts/xlsx_to_json.py file1.xlsx file2.xlsx
    python scripts/xlsx_to_json.py --out data/kumu_blueprint.json data/*.xlsx
"""

import argparse
import json
import sys
from glob import glob
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required — run: pip install openpyxl")


HEADER_MAP = {
    "Task ID":                          "id",
    "Tier":                             "tier",
    "Category":                         "category",
    "Sub-heading":                      "subheading",
    "Source Sentence":                  "source_sentence",
    "Task":                             "task",
    "Lead":                             "lead",
    "Track":                            "track",
    "Recommendation Only":              "rec_only",
    "Connection Type":                  "conn_type",
    "Connected Task IDs":               "conn_ids",
    "Connection Pending Confirmation":  "pending",
    "Notes":                            "notes",
}


def cell(v) -> str:
    return str(v).strip() if v is not None else ""


def yesno(v) -> bool:
    return cell(v).lower() in {"yes", "y", "true", "t", "1", "x"}


def normalise_lead(raw: str) -> str:
    """Reduce compound lead strings to one of four canonical buckets."""
    if not raw:
        return "tbd"
    first = raw.split("/")[0].strip().lower().replace("-", "_").replace(" ", "_")
    if first in {"consultant"}:           return "consultant"
    if first in {"client", "facilities"}: return "client"   # facilities = client-side
    if first in {"third_party"}:          return "third_party"
    return "tbd"


def normalise_tier(raw: str) -> str:
    return raw.strip().lower() if raw else ""


def normalise_conn_type(raw: str) -> str:
    s = raw.strip().lower()
    if s in {"dependency", "synergy"}:
        return s
    return ""   # blank or "—" → no connection


def split_ids(raw: str) -> list[str]:
    if not raw:
        return []
    parts = raw.replace(";", ",").split(",")
    return [p.strip() for p in parts if p.strip()]


def read_workbook(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Task Database"] if "Task Database" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [HEADER_MAP.get(str(h or "").strip(), str(h or "").strip()) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None for v in r):
            continue
        d = dict(zip(headers, r))
        if not cell(d.get("id")):
            continue
        out.append(d)
    return out


def convert(xlsx_paths: list[Path], out_path: Path) -> None:
    all_rows: list[dict] = []
    for p in xlsx_paths:
        rows = read_workbook(p)
        print(f"  · {p.name}: {len(rows)} rows")
        all_rows.extend(rows)

    elements: list[dict] = []
    seen_ids: set[str] = set()
    for r in all_rows:
        tid = cell(r.get("id"))
        if tid in seen_ids:
            print(f"  ⚠  duplicate id '{tid}' — skipping", file=sys.stderr)
            continue
        seen_ids.add(tid)

        lead_raw = cell(r.get("lead"))
        elements.append({
            "id":    tid,
            "label": cell(r.get("task")) or tid,
            "attributes": {
                "tier":                normalise_tier(cell(r.get("tier"))),
                "category":            cell(r.get("category")),
                "subheading":          cell(r.get("subheading")),
                "source_sentence":     cell(r.get("source_sentence")),
                "lead":                normalise_lead(lead_raw),
                "lead_label":          lead_raw,
                "track":               cell(r.get("track")),
                "recommendation_only": yesno(r.get("rec_only")),
                "notes":               cell(r.get("notes")),
            },
        })

    connections: list[dict] = []
    seen_synergy: set[tuple[str, str]] = set()
    dangling: list[tuple[str, str]] = []

    for r in all_rows:
        src_row_id = cell(r.get("id"))
        ctype      = normalise_conn_type(cell(r.get("conn_type")))
        if not ctype:
            continue
        targets = split_ids(cell(r.get("conn_ids")))
        if not targets:
            continue
        pending = yesno(r.get("pending"))
        status  = "pending" if pending else "confirmed"

        for other_id in targets:
            if other_id not in seen_ids:
                dangling.append((src_row_id, other_id))
                continue

            if ctype == "dependency":
                # row depends on other → arrow flows other → row (prereq → dependent)
                connections.append({
                    "from":  other_id,
                    "to":    src_row_id,
                    "label": "",
                    "attributes": {"type": "dependency", "status": status},
                })
            else:  # synergy — undirected, dedupe pair
                pair = tuple(sorted([src_row_id, other_id]))
                if pair in seen_synergy:
                    continue
                seen_synergy.add(pair)
                connections.append({
                    "from":  pair[0],
                    "to":    pair[1],
                    "label": "",
                    "attributes": {"type": "synergy", "status": status},
                })

    if dangling:
        print(f"  ⚠  {len(dangling)} dangling refs (target id not in dataset):", file=sys.stderr)
        for src, tgt in dangling[:10]:
            print(f"       {src} → {tgt}", file=sys.stderr)
        if len(dangling) > 10:
            print(f"       … +{len(dangling) - 10} more", file=sys.stderr)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"elements": elements, "connections": connections}
    out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False))

    print(f"\n  ✓  {len(elements)} tasks, {len(connections)} connections → {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert audit XLSX → kumu_blueprint.json")
    parser.add_argument("xlsx", nargs="*", help="XLSX file(s); defaults to data/*.xlsx")
    parser.add_argument("--out", type=Path, default=Path("data/kumu_blueprint.json"))
    args = parser.parse_args()

    if args.xlsx:
        paths = [Path(p) for pattern in args.xlsx for p in glob(pattern)] or [Path(p) for p in args.xlsx]
    else:
        paths = sorted(Path("data").glob("*.xlsx"))

    paths = [p for p in paths if p.exists()]
    if not paths:
        sys.exit("No XLSX files found.")

    print(f"Reading {len(paths)} workbook(s):")
    convert(paths, args.out)


if __name__ == "__main__":
    main()
